import csv, json, pathlib
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont, ImageFilter
import random
random.seed(7)
out=pathlib.Path("corpus/synth")
# DOCX
d=Document(); d.add_heading("Acme Orbital: Q3 board report",0)
d.add_heading("Summary",1)
d.add_paragraph("Acme Orbital closed the third quarter with revenue of $4.2 million, up 12 percent on the second quarter. The company is headquartered in Lisbon, and its chief executive is Priya Raman.")
d.add_heading("Revenue by region",1)
t=d.add_table(rows=1,cols=3); t.style="Table Grid"
for i,h in enumerate(["Region","Revenue (USD million)","Change"]): t.rows[0].cells[i].text=h
for r in [("APAC","1.7","+18%"),("EMEA","1.5","+9%"),("Americas","1.0","+6%")]:
    c=t.add_row().cells
    for i,v in enumerate(r): c[i].text=v
d.add_heading("Risks",1)
d.add_paragraph("The main risk is the single supplier of star trackers, Helios Optics, which has quoted a 14 week lead time. The board asked management to qualify a second supplier by the end of the year.")
d.save(out/"q3-board-report.docx")
# PPTX 30 slides
p=Presentation()
for i in range(1,31):
    s=p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text=f"Launch plan: step {i}"
    body=s.placeholders[1].text_frame
    body.text=f"Workstream {i} owner reviews milestone {i}."
    if i==1: body.text="The product is called Kestrel. Public launch is planned for 14 March 2027."
    if i==2: body.text="The marketing budget for the launch is $350,000."
    if i==17: body.text="The beta programme will start with 25 customers in the first week of January."
p.save(out/"launch-plan.pptx")
# XLSX
wb=Workbook(); ws=wb.active; ws.title="Staff"; ws.append(["Team","People","Open roles"])
for r in [("Engineering",48,5),("Sales",27,3),("Operations",31,1),("Finance",9,0),("Support",16,2)]: ws.append(r)
ws.append(["Total","=SUM(B2:B6)","=SUM(C2:C6)"])
wb.save(out/"headcount.xlsx")
# HTML
(out/"faq.html").write_text("<html><head><title>Support FAQ</title></head><body><h1>Support FAQ</h1><h2>Hours</h2><p>Support is available from 8am to 6pm UTC on weekdays.</p><h2>Refunds</h2><p>Hardware can be returned for a refund within 30 days of delivery.</p></body></html>")
# MD, CSV
(out/"notes.md").write_text("# Pilot notes\n\nThe pilot customer is Northwind Dairy. They run 12 ground stations across three countries.\n\n## Open questions\n\n- Who pays for the antenna calibration?\n- Can the pilot be extended to Q2?\n")
with open(out/"tickets.csv","w",newline="") as f:
    w=csv.writer(f); w.writerow(["ticket","customer","topic","status"])
    for r in [(101,"Northwind Dairy","antenna calibration","open"),(102,"Northwind Dairy","billing","closed"),(103,"Verdant Farms","firmware update","open")]: w.writerow(r)
# scanned PDF: two pages as images
font=ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc",34)
pages=[]
texts=[["Acme Orbital Warranty Terms","","The warranty period for all ground station","hardware is 18 months from the date of delivery.","Batteries are covered for 6 months only."],
       ["Exclusions","","Damage caused by lightning is not covered.","Repairs are performed at the Lisbon service centre."]]
for lines in texts:
    im=Image.new("L",(1240,1000),255); dr=ImageDraw.Draw(im); y=80
    for ln in lines: dr.text((80,y),ln,font=font,fill=20); y+=60
    im=im.rotate(0.6,fillcolor=255).filter(ImageFilter.GaussianBlur(0.8))
    pages.append(im.convert("RGB"))
pages[0].save(out/"warranty-scan.pdf",save_all=True,append_images=pages[1:],resolution=150)
qs=[
 ("What was Acme Orbital's Q3 revenue?","4.2 million","q3-board-report.docx"),
 ("Which region had the highest revenue?","APAC","q3-board-report.docx"),
 ("Who is the chief executive?","Priya Raman","q3-board-report.docx"),
 ("Where is the company headquartered?","Lisbon","q3-board-report.docx"),
 ("Who supplies the star trackers and what is the lead time?","Helios","q3-board-report.docx"),
 ("What is the name of the product and when does it launch?","14 March 2027","launch-plan.pptx"),
 ("What is the marketing budget for the launch?","350,000","launch-plan.pptx"),
 ("How many customers are in the beta programme?","25","launch-plan.pptx"),
 ("How many people work in Engineering?","48","headcount.xlsx"),
 ("How many open roles does Sales have?","3","headcount.xlsx"),
 ("What are the support hours?","8am to 6pm","faq.html"),
 ("What is the refund window for hardware?","30 days","faq.html"),
 ("Who is the pilot customer?","Northwind Dairy","notes.md"),
 ("How long is the hardware warranty?","18 months","warranty-scan.pdf"),
 ("Which service centre does repairs?","Lisbon","warranty-scan.pdf"),
 ("What is the capital of Australia?",None,None),
 ("What is Acme Orbital's stock price?",None,None),
 ("Who is the chief financial officer?",None,None),
 ("How many satellites has Acme launched?",None,None),
 ("What is the warranty on software licences?",None,None),
]
json.dump([{"q":q,"expect":e,"source":s} for q,e,s in qs],open("questions.json","w"),indent=1)
print("ok")
