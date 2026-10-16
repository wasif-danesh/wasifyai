# pyright: reportAttributeAccessIssue=false
"""Studio: the formats a notebook makes, from the reply to the file, and the run that makes them."""

import io
import json
import sqlite3

import pytest
from notebooks_rig import SPEC, FakeComponents
from test_notebooks_chat import BUDGET, PEOPLE, Rig, ctx  # noqa: F401  (ctx is a fixture)
from wasifyai_notebooks.retrieval import Hit
from wasifyai_notebooks.studio.context import (
    DEFAULT_QUERY,
    budget_chars,
    gather,
    grounding,
)
from wasifyai_notebooks.studio.formats import FORMATS, BadReply, build, parse_json, prompt_for

REPORT = (
    "# Launch plan\n\n## Budget\nThe marketing budget for the launch is **$350,000**.\n\n"
    "## People\n- Engineering has 48 people\n- 5 roles are open\n\n"
    "| Team | People |\n|---|---|\n| Engineering | 48 |\n| Sales | 27 |\n"
)
CARDS = {
    "title": "Launch facts",
    "cards": [
        {"front": "What is the marketing budget?", "back": "$350,000."},
        {"front": "How many people work in engineering?", "back": "48."},
        {"front": "", "back": "dropped: no front"},
        "not a card",
    ],
}
QUIZ = {
    "title": "Launch quiz",
    "questions": [
        {
            "question": "What is the marketing budget?",
            "options": ["$100,000", "$350,000", "$500,000", "$1 million"],
            "answer": "$350,000",
            "explanation": "The plan sets it at $350,000.",
        },
        {  # an answer that is not one of its options cannot be scored
            "question": "Broken?",
            "options": ["a", "b", "c", "d"],
            "answer": "e",
            "explanation": "",
        },
        {  # three options: dropped
            "question": "Short?",
            "options": ["a", "b", "c"],
            "answer": "a",
            "explanation": "",
        },
    ],
}
MAP = {
    "title": "Launch",
    "nodes": [
        {"label": "Budget", "children": [{"label": "$350,000"}, {"label": "Marketing"}]},
        {
            "label": "People",
            "children": [{"label": "Engineering 48", "children": [{"label": "5 open"}]}],
        },
        {"label": ""},
    ],
}


def test_every_format_has_its_prompts_for_both_tiers():
    for fmt in FORMATS.values():
        for tier in ("compact", "capable"):
            text = prompt_for(fmt, tier, "the budget")
            assert text.strip() and "$focus" not in text and "$options" not in text
            if fmt.key != "guide":  # the guide is the same for everyone
                assert "Focus on: the budget" in text
    # a study guide asks for itself when the person said nothing
    assert "study guide" in prompt_for(FORMATS["study_guide"], "capable", None)
    assert "Focus on:" not in prompt_for(FORMATS["report"], "capable", None)


def test_a_report_keeps_its_markdown_title_and_is_a_markdown_file():
    built = build(FORMATS["report"], f"```markdown\n{REPORT}```")
    assert built.title == "Launch plan" and built.markdown.startswith("# Launch plan")
    assert built.file_ext == "md" and built.file == built.markdown.encode()
    with pytest.raises(BadReply):
        build(FORMATS["report"], "ok")


def test_flashcards_keep_only_whole_cards():
    built = build(FORMATS["flashcards"], "Here you go:\n```json\n" + json.dumps(CARDS) + "\n```")
    assert built.data is not None
    assert built.data["cards"] == [
        {"front": "What is the marketing budget?", "back": "$350,000."},
        {"front": "How many people work in engineering?", "back": "48."},
    ]
    assert "**1. What is the marketing budget?**" in built.markdown
    assert built.file_ext == "json" and json.loads(built.file or b"")["title"] == "Launch facts"
    with pytest.raises(BadReply, match="no usable flashcards"):
        build(FORMATS["flashcards"], json.dumps({"title": "x", "cards": []}))
    with pytest.raises(BadReply, match="not valid JSON"):
        build(FORMATS["flashcards"], "no json here")


def test_a_quiz_keeps_only_questions_it_can_score():
    built = build(FORMATS["quiz"], json.dumps(QUIZ))
    assert built.data is not None
    (q,) = built.data["questions"]
    assert q["answer_index"] == 1 and q["options"][1] == "$350,000"
    assert "_Answer: $350,000_" in built.markdown
    with pytest.raises(BadReply, match="four options"):
        build(FORMATS["quiz"], json.dumps({"title": "t", "questions": [QUIZ["questions"][2]]}))


def test_a_mind_map_is_an_outline_three_levels_deep_at_most():
    built = build(FORMATS["mindmap"], json.dumps(MAP))
    assert built.markdown == (
        "# Launch\n- Budget\n  - $350,000\n  - Marketing\n"
        "- People\n  - Engineering 48\n    - 5 open"
    )
    assert built.file is None and built.data is None
    with pytest.raises(BadReply):
        build(FORMATS["mindmap"], json.dumps({"title": "x", "nodes": []}))


def test_word_and_pdf_are_built_from_the_same_markdown():
    import docx

    word = build(FORMATS["word"], REPORT)
    assert word.file_ext == "docx" and word.file is not None
    document = docx.Document(io.BytesIO(word.file))
    texts = [p.text for p in document.paragraphs]
    assert "Launch plan" in texts and any("$350,000" in t for t in texts)
    assert len(document.tables) == 1 and document.tables[0].rows[1].cells[1].text == "48"
    styles = [
        str(p.style.name)
        for p in document.paragraphs
        if p.style is not None and "Engineering has" in p.text
    ]
    assert styles == ["List Bullet"]

    pdf = build(FORMATS["pdf"], REPORT)
    assert pdf.file_ext == "pdf" and pdf.file is not None
    assert pdf.file.startswith(b"%PDF-") and b"%%EOF" in pdf.file[-64:]
    assert len(pdf.file) > 1500


def test_a_document_never_runs_what_it_contains_and_links_stay_safe():
    import docx

    risky = (
        "# Notes\n\nSee [the site](https://example.com/a) and [bad](javascript:alert(1)). "
        "<script>alert(1)</script> and a long enough body to count as a document.\n"
    )
    word = build(FORMATS["word"], risky)
    assert word.file is not None
    doc = docx.Document(io.BytesIO(word.file))
    body = " ".join(p.text for p in doc.paragraphs)
    assert "<script>" in body  # raw HTML stays text
    rels = [r.target_ref for r in doc.part.rels.values() if r.is_external]
    assert rels == ["https://example.com/a"]  # the javascript: link was not made a link
    assert build(FORMATS["pdf"], risky).file is not None


def test_json_is_found_inside_a_fence_or_a_sentence():
    assert parse_json('Sure!\n```json\n{"a": 1}\n```\nDone.') == {"a": 1}
    assert parse_json('The answer is {"a": {"b": 2}} as asked.') == {"a": {"b": 2}}
    with pytest.raises(BadReply):
        parse_json("[1, 2]")


# ---- cutting the sources to the window ---------------------------------------------------------


def _db(tmp_path, docs):
    db = sqlite3.connect(tmp_path / "t.db")
    db.executescript(
        "CREATE TABLE documents(id TEXT, notebook_id TEXT, title TEXT, content TEXT, status TEXT);"
    )
    for i, (title, content) in enumerate(docs):
        db.execute(
            "INSERT INTO documents VALUES (?,?,?,?, 'ready')", (f"d{i}", "n", title, content)
        )
    return db


def test_a_selection_that_fits_is_sent_whole(tmp_path):
    db = _db(tmp_path, [("a.txt", "alpha " * 50), ("b.txt", "beta " * 50)])
    out = gather(db, "n", ["d0", "d1", "d9"], budget=10_000)
    assert [s.title for s in out] == ["a.txt", "b.txt"]
    assert out[0].content == "alpha " * 50
    text = grounding(out)
    assert text.startswith("# a.txt\n\nalpha") and "# b.txt" in text
    assert grounding([]) == "No sources were provided."


def test_a_selection_that_does_not_fit_gets_an_even_share_each(tmp_path):
    db = _db(tmp_path, [("a.txt", "a" * 9000), ("b.txt", "b" * 1000), ("c.txt", "c" * 9000)])
    out = gather(db, "n", ["d0", "d1", "d2"], budget=6000)
    assert [len(s.content) for s in out] == [
        2500,
        1000,
        2500,
    ]  # the short one gives its unused part
    assert sum(len(s.content) for s in out) <= 6000


def test_many_sources_fall_back_to_the_best_passages(tmp_path, monkeypatch):
    docs = [(f"d{i}.txt", f"text {i} " * 400) for i in range(6)]
    db = _db(tmp_path, docs)
    seen = {}

    def fake_retrieve(_db, nid, query, vec, *, weight, top_k, document_ids):
        seen.update(query=query, top_k=top_k, ids=list(document_ids))
        return [Hit(5, "d3", "d3.txt", "the passage that matters", 1, 1, 0.9),
                Hit(4, "d1", "d1.txt", "another passage", 1, 1, 0.8)]  # fmt: skip

    monkeypatch.setattr("wasifyai_notebooks.studio.context.retrieve", fake_retrieve)
    out = gather(
        db, "n", [f"d{i}" for i in range(6)], budget=6000, embed=lambda t: [0.0] * 4, query=None
    )
    assert seen["query"] == DEFAULT_QUERY and seen["top_k"] == 4
    assert [(s.document_id, s.content) for s in out] == [
        ("d1", "another passage"),
        ("d3", "the passage that matters"),
    ]  # only sources holding a passage, in the order selected
    assert budget_chars(8192) > 15_000 and budget_chars(1024) == 4000


# ---- a run that makes an output ------------------------------------------------------------------


async def studio(rig, fmt, prompt="", document_ids=None):
    graph_input = {"notebook_id": rig.nb["id"], "studio": {"format": fmt, "prompt": prompt}}
    if document_ids is not None:
        graph_input["document_ids"] = document_ids
    config = {"configurable": {"thread_id": f"s-{fmt}-{rig._n}"}}
    rig._n += 1
    final, events = {}, []
    graph = rig.graph
    async for mode, chunk in graph.astream(graph_input, config, stream_mode=["custom", "values"]):
        if mode == "custom":
            events.append(chunk)
        else:
            final = chunk
    return final, events


async def test_a_report_is_written_saved_with_its_file_and_listed(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[REPORT])
    await rig.add("budget.txt", BUDGET)
    await rig.add("hr.txt", PEOPLE)
    state, events = await studio(rig, "report", "the launch")
    assert state["status"] == "done" and state["title"] == "Launch plan"
    art = rig.store.get_artifact("t1", "u1", rig.nb["id"], state["artifact_id"], content=True)
    assert art["status"] == "ready" and art["format"] == "report"
    assert art["body"].startswith("# Launch plan") and art["prompt"] == "the launch"
    assert art["file_name"] == "launch-plan.md" and len(art["source_ids"]) == 2
    saved = rig.runtime.artifacts_dir / rig.nb["id"] / art["id"] / "launch-plan.md"
    assert saved.read_text().startswith("# Launch plan")
    assert [e["node"] for e in events if e["type"] == "node"] == ["read", "write"]
    # the sources were put in front of the model under the prompt for its tier
    sent = rig.model.prompts[0]
    assert "# budget.txt" in str(sent) and "$350,000" in str(sent)
    assert [a["id"] for a in rig.store.list_artifacts("t1", "u1", rig.nb["id"])] == [art["id"]]


async def test_flashcards_come_back_as_data_for_the_viewer(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[json.dumps(CARDS)])
    await rig.add("budget.txt", BUDGET)
    state, _ = await studio(rig, "flashcards")
    art = rig.store.get_artifact("t1", "u1", rig.nb["id"], state["artifact_id"], content=True)
    assert art["data"]["cards"][0] == {
        "front": "What is the marketing budget?",
        "back": "$350,000.",
    }
    assert art["file_name"].endswith(".json")


async def test_a_bad_reply_is_tried_once_more_then_fails_in_words(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=["not json at all", "still not json"])
    await rig.add("budget.txt", BUDGET)
    state, _ = await studio(rig, "quiz")
    art = rig.store.get_artifact("t1", "u1", rig.nb["id"], state["artifact_id"])
    assert art["status"] == "failed" and "could not be used" in art["error"]
    assert len(rig.model.prompts) == 2
    assert "previous reply could not be used" in str(rig.model.prompts[1])
    # the second attempt can succeed
    rig2 = Rig(tmp_path / "again", [{"text": "x"}], replies=["oops", json.dumps(QUIZ)])
    await rig2.add("budget.txt", BUDGET)
    state2, _ = await studio(rig2, "quiz")
    assert (
        rig2.store.get_artifact("t1", "u1", rig2.nb["id"], state2["artifact_id"])["status"]
        == "ready"
    )


async def test_the_ticked_sources_are_the_only_ones_used(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[REPORT])
    a = await rig.add("budget.txt", BUDGET)
    await rig.add("hr.txt", PEOPLE)
    state, _ = await studio(rig, "word", document_ids=[a["id"]])
    art = rig.store.get_artifact("t1", "u1", rig.nb["id"], state["artifact_id"])
    assert art["source_ids"] == [a["id"]] and art["file_name"] == "launch-plan.docx"
    sent = str(rig.model.prompts[0])
    assert "budget.txt" in sent and "hr.txt" not in sent


async def test_nothing_to_make_it_from_is_refused_before_the_model(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[REPORT])
    state, _ = await studio(rig, "report")
    assert state["refusal"]["code"] == "no_sources"
    await rig.add("budget.txt", BUDGET)
    state, _ = await studio(rig, "nope")
    assert state["refusal"]["code"] == "invalid_request"
    state, _ = await studio(rig, "report", document_ids=["gone"])
    assert state["refusal"]["code"] == "no_scope"
    assert rig.model.prompts == [] and rig.store.list_artifacts("t1", "u1", rig.nb["id"]) == []


async def test_an_output_stopped_meanwhile_is_not_overwritten(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[REPORT])
    await rig.add("budget.txt", BUDGET)
    original = rig.store.finish_artifact

    def stopped(t, u, nid, aid, **kw):
        rig.store.cancel_artifact(t, u, nid, aid)  # the person stops it as it finishes
        return original(t, u, nid, aid, **kw)

    rig.store.finish_artifact = stopped  # type: ignore[method-assign]
    state, _ = await studio(rig, "report")
    art = rig.store.get_artifact("t1", "u1", rig.nb["id"], state["artifact_id"])
    assert art["status"] == "cancelled"
    assert not (rig.runtime.artifacts_dir / rig.nb["id"] / art["id"]).exists()


def test_outputs_are_owned_and_removed_with_their_notebook(tmp_path):
    from wasifyai_notebooks.store import NotebookStore, NotFound

    store = NotebookStore(tmp_path, dimension=SPEC.dimension, embedder_id=SPEC.id)
    nb = store.create_notebook("t1", "u1", "n")
    a = store.create_artifact("t1", "u1", nb["id"], fmt="report", prompt="", source_ids=["d"])
    assert a["status"] == "processing"
    with pytest.raises(NotFound):
        store.get_artifact("t1", "u2", nb["id"], a["id"])
    assert store.reset_interrupted() == 0
    got = store.get_artifact("t1", "u1", nb["id"], a["id"])
    assert got["status"] == "failed" and "interrupted" in got["error"]
    store.delete_notebook("t1", "u1", nb["id"])
    assert FakeComponents  # keep the import used


# ---- E5: slides, tables, audio script, infographic, the guide ----------------------------------

DECK = {
    "title": "Launch",
    "slides": [
        {
            "title": "Budget",
            "bullets": ["$350,000 for marketing", "Launch on 14 March 2027"],
            "notes": "Board approved.",
        },
        {"title": "No bullets", "bullets": [], "notes": ""},  # dropped: nothing to show
        {"title": "People", "bullets": ["48 in engineering"], "notes": ""},
    ],
}
TABLE = {
    "title": "Headcount",
    "columns": ["Team", "People", "Note"],
    "rows": [["Engineering", "48", "=SUM(A1:A9)"], ["Sales", "1,250", ""], [], ["", "", ""]],
}
SCRIPT = {
    "title": "Launch overview",
    "turns": [
        {"speaker": 1, "text": "Today: what is in the launch plan?"},
        {"speaker": 2, "text": "A March launch with a $350,000 budget."},
        {"speaker": 1, "text": "And who builds it?"},
        {"speaker": 2, "text": "Forty-eight engineers."},
        {"speaker": 3, "text": "dropped: not a host"},
    ],
}
INFO = {
    "title": "Launch at a glance",
    "summary": "A March launch with a small beta.",
    "sections": [
        {"label": "Launch", "value": "14 March 2027", "detail": "Public"},
        {"label": "Budget", "value": "$350,000", "detail": "Marketing"},
        {"label": "", "value": "dropped", "detail": ""},
    ],
}
GUIDE = {
    "summary": "These sources describe a launch: its date, its budget and its team.",
    "topics": ["Launch date", "Budget", "Team"],
    "questions": ["When is the launch?", "Not a question", "What is the budget?"],
}


def test_slides_become_a_deck_with_notes_and_nothing_empty():
    from pptx import Presentation

    built = build(FORMATS["slides"], json.dumps(DECK))
    assert built.file_ext == "pptx" and built.data is not None
    assert [s["title"] for s in built.data["slides"]] == ["Budget", "People"]
    deck = Presentation(io.BytesIO(built.file or b""))
    assert len(deck.slides) == 3  # a title slide and two
    texts = [sh.text_frame.text for sh in deck.slides[1].shapes if sh.has_text_frame]
    assert "Budget" in texts and any("$350,000" in t for t in texts)
    notes = deck.slides[1].notes_slide.notes_text_frame
    assert notes is not None and notes.text == "Board approved."
    with pytest.raises(BadReply, match="no usable slides"):
        build(FORMATS["slides"], json.dumps({"title": "x", "slides": [DECK["slides"][1]]}))


def test_a_table_becomes_a_workbook_of_values_never_formulas():
    import openpyxl

    built = build(FORMATS["datatable"], json.dumps(TABLE))
    assert built.data is not None and built.data["rows"] == [
        ["Engineering", "48", "=SUM(A1:A9)"],
        ["Sales", "1,250", ""],
    ]  # the empty rows are gone, short rows are padded
    book = openpyxl.load_workbook(io.BytesIO(built.file or b""))
    sheet = book.active
    assert sheet is not None
    assert [c.value for c in sheet[1]] == ["Team", "People", "Note"]
    assert sheet["B2"].value == 48 and sheet["B3"].value == 1250  # numbers are numbers
    assert (
        sheet["C2"].data_type == "s" and sheet["C2"].value == "=SUM(A1:A9)"
    )  # text, not a formula
    assert built.markdown.splitlines()[2].startswith("| Team | People | Note |")
    with pytest.raises(BadReply):
        build(FORMATS["datatable"], json.dumps({"title": "t", "columns": [], "rows": [["a"]]}))
    with pytest.raises(BadReply, match="no rows"):
        build(FORMATS["datatable"], json.dumps({"title": "t", "columns": ["a"], "rows": [[]]}))


def test_an_audio_overview_is_a_script_between_two_named_hosts():
    built = build(FORMATS["audio"], json.dumps(SCRIPT))
    assert built.data is not None
    assert [t["host"] for t in built.data["turns"]] == ["Alex", "Sam", "Alex", "Sam"]
    assert built.data["audio"] is False  # the voices need an engine this build does not have
    assert (
        built.markdown.startswith("# Launch overview")
        and "**Sam:** A March launch" in built.markdown
    )
    assert built.file_ext == "md"
    with pytest.raises(BadReply, match="too few"):
        build(FORMATS["audio"], json.dumps({"title": "t", "turns": SCRIPT["turns"][:2]}))


def test_an_infographic_is_a_brief_and_a_prompt_within_the_image_limit():
    from wasifyai_notebooks.studio.formats import IMAGE_PROMPT_CHARS, infographic_prompt

    built = build(FORMATS["infographic"], json.dumps(INFO))
    assert built.data is not None and len(built.data["sections"]) == 2 and built.file is None
    prompt = infographic_prompt(built.data)
    assert "Launch at a glance" in prompt and "Budget: $350,000" in prompt
    many = {
        "title": "T",
        "sections": [{"label": f"Label {i}", "value": "v" * 40} for i in range(8)],
    }
    assert len(infographic_prompt(many)) <= IMAGE_PROMPT_CHARS


def test_a_guide_keeps_only_real_questions():
    built = build(FORMATS["guide"], json.dumps(GUIDE))
    assert built.data is not None
    assert built.data["questions"] == ["When is the launch?", "What is the budget?"]
    assert built.data["topics"] == ["Launch date", "Budget", "Team"]
    with pytest.raises(BadReply):
        build(FORMATS["guide"], json.dumps({"summary": "", "topics": [], "questions": []}))


class Painter:
    """A tool host that draws like the image product: events of a child run."""

    def __init__(self, events):
        self.events = events
        self.inputs = []

    async def run_product(self, product, graph_input):
        self.inputs.append((product, graph_input))
        for e in self.events:
            yield e


async def test_the_guide_is_saved_and_goes_stale_when_sources_change(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[json.dumps(GUIDE)])
    await rig.add("budget.txt", BUDGET)
    before = rig.store.get_guide("t1", "u1", rig.nb["id"])
    assert before == {"guide": None, "stale": False, "ready": True}
    state, _ = await studio(rig, "guide")
    assert state["status"] == "done" and "artifact_id" not in state
    got = rig.store.get_guide("t1", "u1", rig.nb["id"])
    assert got["guide"]["questions"] == ["When is the launch?", "What is the budget?"]
    assert got["stale"] is False
    assert rig.store.list_artifacts("t1", "u1", rig.nb["id"]) == []  # not an output
    await rig.add("hr.txt", PEOPLE)
    assert rig.store.get_guide("t1", "u1", rig.nb["id"])["stale"] is True


async def test_an_infographic_asks_the_image_product_for_its_picture(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[json.dumps(INFO)])
    await rig.add("budget.txt", BUDGET)
    host = Painter([
        {"event": "job_progress", "status": "running"},
        {"event": "done", "outputs": {"status": "done", "image_id": "img-1",
                                      "width": 864, "height": 1152}},
    ])  # fmt: skip
    rig.caps.extras["tool_host"] = host
    state, events = await studio(rig, "infographic")
    art = rig.store.get_artifact("t1", "u1", rig.nb["id"], state["artifact_id"], content=True)
    assert art["status"] == "ready" and art["data"]["image_id"] == "img-1"
    (product, sent) = host.inputs[0]
    assert product == "wasifyai-image" and sent["size"] == "portrait" and len(sent["prompt"]) <= 500
    assert [e["node"] for e in events if e["type"] == "node"] == ["read", "write", "paint"]


async def test_an_infographic_says_why_when_no_picture_could_be_made(tmp_path, ctx):  # noqa: F811
    rig = Rig(tmp_path, [{"text": "x"}], replies=[json.dumps(INFO)])
    await rig.add("budget.txt", BUDGET)
    state, _ = await studio(rig, "infographic")  # no tool host here
    art = rig.store.get_artifact("t1", "u1", rig.nb["id"], state["artifact_id"])
    assert art["status"] == "failed" and "not available here" in art["error"]
    rig.caps.extras["tool_host"] = Painter(
        [{"event": "error", "message": "The image engine is not installed."}]
    )
    state2, _ = await studio(rig, "infographic")
    art2 = rig.store.get_artifact("t1", "u1", rig.nb["id"], state2["artifact_id"])
    assert art2["status"] == "failed" and "image engine is not installed" in art2["error"]
