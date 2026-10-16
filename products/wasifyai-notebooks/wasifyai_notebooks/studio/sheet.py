# SPDX-License-Identifier: Apache-2.0
"""A workbook from a validated table (the model writes JSON, this renders the file).

Values only: nothing a model wrote becomes a formula."""

import re
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_COLUMNS = 12
MAX_ROWS = 200
MAX_CELL = 500
_NUMBER = re.compile(r"^-?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")


def cell_value(text: Any) -> Any:
    """A number when the text is one ("1,250", "48", "3.5"), else the text; empty is empty."""
    s = str(text if text is not None else "").strip()[:MAX_CELL]
    if not s:
        return None
    if _NUMBER.match(s):
        plain = s.replace(",", "")
        return float(plain) if "." in plain else int(plain)
    return s


def build_workbook(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = (title or "Table")[:31].replace("/", "-").replace("\\", "-") or "Table"
    sheet.append([str(c)[:MAX_CELL] for c in columns[:MAX_COLUMNS]])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows[:MAX_ROWS]:
        sheet.append([cell_value(v) for v in row[:MAX_COLUMNS]])
    # a cell that starts with `=` is text: the file must not calculate what a model wrote
    for line in sheet.iter_rows(min_row=1):
        for cell in line:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    for i, column in enumerate(columns[:MAX_COLUMNS], start=1):
        longest = max(
            [len(str(column))] + [len(str(r[i - 1])) for r in rows[:MAX_ROWS] if len(r) >= i]
        )
        sheet.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
    out = BytesIO()
    book.save(out)
    return out.getvalue()
