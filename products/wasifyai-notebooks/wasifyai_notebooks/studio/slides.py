# pyright: reportAttributeAccessIssue=false
# SPDX-License-Identifier: Apache-2.0
"""A slide deck from a validated spec (the model writes JSON, this renders the file: ADR-0010).

python-pptx's default template: a title slide, then one "Title and Content" slide per entry with its
bullets and speaker notes. Nothing the model wrote is run.
"""

from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.util import Pt

MAX_SLIDES = 14
MAX_BULLETS = 6
MAX_TITLE = 120
MAX_BULLET = 220


def build_deck(title: str, slides: list[dict[str, Any]]) -> bytes:
    deck = Presentation()
    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title[:MAX_TITLE]  # type: ignore[union-attr]
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = "Made from your sources with Wasify AI"
    for spec in slides[:MAX_SLIDES]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(spec["title"])[:MAX_TITLE]  # type: ignore[union-attr]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        bullets = [str(b)[:MAX_BULLET] for b in spec["bullets"][:MAX_BULLETS]]
        for i, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(24)
        notes = str(spec.get("notes", "")).strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes  # type: ignore[union-attr]
    out = BytesIO()
    deck.save(out)
    return out.getvalue()
